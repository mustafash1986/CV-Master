"""
Excel Job Tracker module.
Creates, formats, and updates the job applications Excel workbook.
"""
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from cv_servant.config import EXCEL_TRACKER_PATH

COLUMNS = [
    "Job ID",
    "Date Detected",
    "Company Name",
    "Job Title",
    "Country",
    "City",
    "Application Method",
    "Contact / Link",
    "Visa Sponsorship",
    "Status",
    "Applied Date",
    "Response Date",
    "Folder Path",
    "Notes"
]

STATUS_COLORS = {
    "Pending Approval": "FFF3CD",   # Soft Yellow
    "Applied / Sent": "D1E7DD",     # Soft Green
    "Interview": "CFF4FC",          # Soft Cyan
    "Under Review": "E2E3E5",       # Soft Grey
    "Rejected": "F8D7DA",           # Soft Red
}


class ExcelTracker:
    def __init__(self, file_path: Path = EXCEL_TRACKER_PATH):
        self.file_path = Path(file_path)
        self._ensure_workbook_exists()

    def _ensure_workbook_exists(self):
        """Create and style workbook if it doesn't exist."""
        if not self.file_path.exists():
            self.file_path.parent.mkdir(parents=True, exist_ok=True)
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Job Applications"

            # Header Style
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1B4965", end_color="1B4965", fill_type="solid")
            thin_border = Border(
                left=Side(style='thin', color='CCCCCC'),
                right=Side(style='thin', color='CCCCCC'),
                top=Side(style='thin', color='CCCCCC'),
                bottom=Side(style='thin', color='CCCCCC')
            )

            ws.append(COLUMNS)

            for col_num, col_name in enumerate(COLUMNS, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin_border
                ws.column_dimensions[get_column_letter(col_num)].width = max(len(col_name) + 4, 15)

            ws.row_dimensions[1].height = 28
            wb.save(str(self.file_path))

    def add_job(self, job_data: Dict[str, Any]) -> str:
        """Add a newly detected/prepared job to the tracker."""
        wb = openpyxl.load_workbook(str(self.file_path))
        ws = wb.active

        job_id = f"JOB-{datetime.now().strftime('%Y%m%d%H%M%S')}"
        today_str = datetime.now().strftime("%Y-%m-%d")

        row_data = [
            job_id,
            today_str,
            job_data.get("company_name", "N/A"),
            job_data.get("job_title", "N/A"),
            job_data.get("country", "N/A"),
            job_data.get("city", ""),
            job_data.get("application_method", "EMAIL"),
            job_data.get("application_email") or job_data.get("job_url", ""),
            job_data.get("visa_sponsorship", "Not Mentioned"),
            job_data.get("status", "Pending Approval"),
            job_data.get("applied_date", ""),
            "",  # Response Date
            str(job_data.get("folder_path", "")),
            job_data.get("sponsorship_notes", "")
        ]

        ws.append(row_data)
        new_row_idx = ws.max_row

        # Style newly added row
        thin_border = Border(
            left=Side(style='thin', color='E0E0E0'),
            right=Side(style='thin', color='E0E0E0'),
            top=Side(style='thin', color='E0E0E0'),
            bottom=Side(style='thin', color='E0E0E0')
        )
        row_font = Font(name="Calibri", size=10)

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=new_row_idx, column=col_idx)
            cell.font = row_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

            # Status cell color
            if col_idx == 10:  # Status
                status_val = str(cell.value)
                if status_val in STATUS_COLORS:
                    cell.fill = PatternFill(start_color=STATUS_COLORS[status_val], end_color=STATUS_COLORS[status_val], fill_type="solid")
                cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[new_row_idx].height = 22
        wb.save(str(self.file_path))
        return job_id

    def update_job_status(self, job_id: str, new_status: str, response_date: Optional[str] = None, notes: Optional[str] = None) -> bool:
        """Update job status (e.g. Sent, Interview, Rejected)."""
        wb = openpyxl.load_workbook(str(self.file_path))
        ws = wb.active

        found = False
        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row=row, column=1).value).strip() == job_id.strip():
                status_cell = ws.cell(row=row, column=10)
                status_cell.value = new_status
                if new_status in STATUS_COLORS:
                    status_cell.fill = PatternFill(start_color=STATUS_COLORS[new_status], end_color=STATUS_COLORS[new_status], fill_type="solid")

                if new_status == "Applied / Sent":
                    ws.cell(row=row, column=11).value = datetime.now().strftime("%Y-%m-%d %H:%M")

                if response_date:
                    ws.cell(row=row, column=12).value = response_date

                if notes:
                    current_notes = ws.cell(row=row, column=14).value or ""
                    ws.cell(row=row, column=14).value = f"{current_notes} | {notes}".strip(" | ")

                found = True
                break

        if found:
            wb.save(str(self.file_path))
        return found

    def get_all_jobs(self) -> List[Dict[str, Any]]:
        """Return all tracked jobs as dictionaries."""
        wb = openpyxl.load_workbook(str(self.file_path))
        ws = wb.active
        jobs = []
        for row in range(2, ws.max_row + 1):
            val_id = ws.cell(row=row, column=1).value
            if not val_id:
                continue
            item = {
                "job_id": val_id,
                "date_detected": ws.cell(row=row, column=2).value,
                "company_name": ws.cell(row=row, column=3).value,
                "job_title": ws.cell(row=row, column=4).value,
                "country": ws.cell(row=row, column=5).value,
                "city": ws.cell(row=row, column=6).value,
                "application_method": ws.cell(row=row, column=7).value,
                "contact": ws.cell(row=row, column=8).value,
                "visa_sponsorship": ws.cell(row=row, column=9).value,
                "status": ws.cell(row=row, column=10).value,
                "applied_date": ws.cell(row=row, column=11).value,
                "response_date": ws.cell(row=row, column=12).value,
                "folder_path": ws.cell(row=row, column=13).value,
                "notes": ws.cell(row=row, column=14).value,
            }
            jobs.append(item)
        return jobs
